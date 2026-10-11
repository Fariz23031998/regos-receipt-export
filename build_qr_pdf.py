"""Build a PDF from an exported receipt workbook.

Each page has the item name and a QR code of its marking.

Usage:
    python build_qr_pdf.py check-1-0751298.xlsx
    python build_qr_pdf.py check-1-0751298.xlsx -o labels.pdf
"""

import argparse
import sys
from io import BytesIO
from pathlib import Path

import fitz
import qrcode
from openpyxl import load_workbook

FONT = Path(r"C:\Windows\Fonts\arial.ttf")
FONT_BOLD = Path(r"C:\Windows\Fonts\arialbd.ttf")


def qr_png(data: str) -> bytes:
    code = qrcode.QRCode(
        error_correction=qrcode.constants.ERROR_CORRECT_H,
        box_size=16,
        border=4,
    )
    code.add_data(data)
    code.make(fit=True)
    image = code.make_image(fill_color="black", back_color="white").convert("RGB")
    buffer = BytesIO()
    image.save(buffer, format="PNG")
    return buffer.getvalue()


def wrap_text(text: str, font: fitz.Font, fontsize: float, max_width: float) -> list[str]:
    words = text.split()
    lines: list[str] = []
    current = ""
    for word in words:
        trial = word if not current else f"{current} {word}"
        if font.text_length(trial, fontsize=fontsize) <= max_width:
            current = trial
        else:
            if current:
                lines.append(current)
            current = word
    if current:
        lines.append(current)
    return lines or [text]


def draw_centered(
    page: fitz.Page,
    lines: list[str],
    top: float,
    font: fitz.Font,
    fontsize: float,
    color: tuple[float, float, float],
    leading: float,
) -> None:
    writer = fitz.TextWriter(page.rect)
    y = top + fontsize
    for line in lines:
        width = font.text_length(line, fontsize=fontsize)
        x = (page.rect.width - width) / 2
        writer.append((x, y), line, font=font, fontsize=fontsize)
        y += leading
    writer.write_text(page, color=color)


def column_index(headers: list[str], name: str) -> int:
    wanted = name.casefold()
    for index, header in enumerate(headers):
        if header.casefold() == wanted:
            return index
    raise SystemExit(f"Excel is missing a '{name}' column. Found: {', '.join(headers) or '(empty)'}")


def read_rows(path: Path) -> list[tuple[str, str]]:
    workbook = load_workbook(path, read_only=True, data_only=True)
    sheet = workbook.active
    rows = sheet.iter_rows(values_only=True)
    header_row = next(rows, None)
    if header_row is None:
        raise SystemExit(f"Excel file is empty: {path}")

    headers = ["" if cell is None else str(cell).strip() for cell in header_row]
    name_index = column_index(headers, "Name")
    marking_index = column_index(headers, "Marking")

    items: list[tuple[str, str]] = []
    skipped = 0
    for row in rows:
        if row is None:
            continue
        name = "" if row[name_index] is None else str(row[name_index]).strip()
        marking = "" if row[marking_index] is None else str(row[marking_index]).strip()
        if not marking:
            skipped += 1
            continue
        items.append((name or "—", marking))

    workbook.close()
    if not items:
        raise SystemExit(f"No markings found in {path}. Skipped {skipped} row(s).")
    return items


def build_pdf(items: list[tuple[str, str]], output: Path) -> None:
    if not FONT.is_file() or not FONT_BOLD.is_file():
        raise SystemExit(f"Arial font not found in {FONT.parent}")

    width, height = fitz.paper_size("a4")
    document = fitz.open()
    font_bold = fitz.Font(fontfile=str(FONT_BOLD))
    font = fitz.Font(fontfile=str(FONT))
    total = len(items)
    qr_size = 430

    for index, (name, marking) in enumerate(items, start=1):
        page = document.new_page(width=width, height=height)
        name_lines = wrap_text(name, font_bold, 26, width - 64)
        draw_centered(page, name_lines, 64, font_bold, 26, (0.1, 0.1, 0.1), 32)
        x0 = (width - qr_size) / 2
        y0 = 180
        page.insert_image(
            fitz.Rect(x0, y0, x0 + qr_size, y0 + qr_size),
            stream=qr_png(marking),
        )
        draw_centered(page, wrap_text(marking, font, 11, width - 72), y0 + qr_size + 28, font, 11, (0.2, 0.2, 0.2), 14)
        draw_centered(page, [f"{index} / {total}"], height - 46, font, 10, (0.45, 0.45, 0.45), 12)

    output.parent.mkdir(parents=True, exist_ok=True)
    document.save(output, deflate=True, garbage=4)
    document.close()


def main() -> None:
    parser = argparse.ArgumentParser(description="Build a QR PDF from a receipt Excel file.")
    parser.add_argument("excel", help="Workbook exported by export_receipt.py")
    parser.add_argument(
        "-o",
        "--output",
        help="PDF path. Defaults to the Excel file name with a .pdf extension.",
    )
    args = parser.parse_args()

    excel = Path(args.excel)
    if not excel.is_file():
        raise SystemExit(f"Excel file not found: {excel}")

    output = Path(args.output) if args.output else excel.with_suffix(".pdf")
    items = read_rows(excel)
    build_pdf(items, output)
    print(f"{len(items)} pages saved to {output}")


if __name__ == "__main__":
    try:
        main()
    except OSError as error:
        print(error, file=sys.stderr)
        raise SystemExit(1)
