"""Export a REGOS receipt to Excel.

Columns: Code, Name, Quantity, Marking.

Usage:
    python export_receipt.py 1-1472645
    python export_receipt.py 1-1472645 -o receipt.xlsx
"""

import argparse
import locale
import sys
from decimal import Decimal
from pathlib import Path

if not hasattr(locale, "resetlocale"):
    locale.resetlocale = lambda *args, **kwargs: None

import os

FIREBIRD_DIR = Path(r"C:\Program Files\Firebird\Firebird_3_0")
if FIREBIRD_DIR.is_dir():
    os.add_dll_directory(str(FIREBIRD_DIR))

import fdb
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

DB_PATH = Path(__file__).resolve().parent / "Regos.FDB"
DB_USER = "SYSDBA"
DB_PASSWORD = "masterkey"

HEADERS = ("Code", "Name", "Quantity", "Marking")

LINE_SQL = """
SELECT i.ITM_CODE, i.ITM_NAME, o.OPR_QUANTITY, o.OPR_LABEL
FROM {table} o
JOIN CTLG_ITM_ITEMS_REF i ON i.ITM_ID = o.OPR_ITEM
WHERE o.OPR_DOCUMENT = CHAR_TO_UUID(?)
ORDER BY o.OPR_ORDER
"""


def connect():
    if not DB_PATH.is_file():
        raise SystemExit(f"Database not found: {DB_PATH}")
    return fdb.connect(
        dsn=f"localhost:{DB_PATH}",
        user=DB_USER,
        password=DB_PASSWORD,
        charset="UTF8",
    )


def find_receipt(cursor, code: str) -> str:
    cursor.execute(
        """
        SELECT UUID_TO_CHAR(SLS_UUID)
        FROM DOCS_DOC_SALES_REF
        WHERE SLS_CODE = ? AND SLS_DELETED = 0
        ORDER BY SLS_DATE DESC
        """,
        (code,),
    )
    row = cursor.fetchone()
    if row is None:
        raise SystemExit(f"Receipt not found: {code}")
    return row[0].strip()


def fetch_lines(cursor, document_uuid: str):
    # Open receipts keep lines in the temp table. Posted receipts keep them
    # in the operations table.
    for table in ("STOCK_OPR_OPERATIONS_TMP", "STOCK_OPR_OPERATIONS_REF"):
        cursor.execute(LINE_SQL.format(table=table), (document_uuid,))
        rows = cursor.fetchall()
        if rows:
            return rows
    return []


def quantity_value(value):
    if value is None:
        return None
    number = Decimal(value)
    if number == number.to_integral_value():
        return int(number)
    return float(number)


def save_excel(path: Path, code: str, rows) -> None:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = code[:31]
    sheet.append(HEADERS)

    header_font = Font(bold=True, color="FFFFFF")
    header_fill = PatternFill("solid", fgColor="1F4E79")
    for column in range(1, len(HEADERS) + 1):
        cell = sheet.cell(1, column)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center")

    for item_code, name, quantity, marking in rows:
        sheet.append([
            int(item_code),
            name,
            quantity_value(quantity),
            marking or "",
        ])

    for row in sheet.iter_rows(min_row=2, max_row=sheet.max_row, min_col=1, max_col=4):
        row[0].alignment = Alignment(horizontal="center")
        row[2].alignment = Alignment(horizontal="center")

    sheet.auto_filter.ref = f"A1:D{sheet.max_row}"
    sheet.freeze_panes = "A2"
    widths = (12, 42, 12, 55)
    for index, width in enumerate(widths, start=1):
        sheet.column_dimensions[get_column_letter(index)].width = width

    path.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(path)


def main() -> None:
    parser = argparse.ArgumentParser(description="Export a REGOS receipt to Excel.")
    parser.add_argument("code", help="Receipt code, for example 1-1472645")
    parser.add_argument(
        "-o",
        "--output",
        help="Excel file path. Defaults to check-<code>.xlsx next to this script.",
    )
    args = parser.parse_args()

    code = args.code.strip()
    output = Path(args.output) if args.output else Path(__file__).resolve().parent / f"check-{code}.xlsx"

    connection = connect()
    try:
        cursor = connection.cursor()
        document_uuid = find_receipt(cursor, code)
        rows = fetch_lines(cursor, document_uuid)
    finally:
        connection.close()

    if not rows:
        raise SystemExit(f"Receipt {code} has no lines.")

    save_excel(output, code, rows)
    print(f"{len(rows)} lines saved to {output}")


if __name__ == "__main__":
    try:
        main()
    except fdb.Error as error:
        print(error, file=sys.stderr)
        raise SystemExit(1)
